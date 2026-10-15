from decimal import Decimal, InvalidOperation
import re
import sys
import unicodedata


def _importar_load_workbook():
    """Carga openpyxl sin depender de instalaciones incompletas de NumPy."""
    ausente = object()
    numpy_previo = sys.modules.get("numpy", ausente)
    numpy_bloqueado = False

    try:
        import numpy as numpy_opcional
    except ImportError:
        pass
    else:
        if not hasattr(numpy_opcional, "short"):
            # openpyxl admite funcionar sin NumPy, pero falla si encuentra un
            # paquete incompleto. Bloquearlo temporalmente hace que use sus
            # tipos numéricos nativos de Python.
            sys.modules["numpy"] = None
            numpy_bloqueado = True

    try:
        from openpyxl import load_workbook as cargar_libro
    finally:
        if numpy_bloqueado:
            if numpy_previo is ausente:
                sys.modules.pop("numpy", None)
            else:
                sys.modules["numpy"] = numpy_previo

    return cargar_libro


load_workbook = _importar_load_workbook()


HOJA_PRECIOS = "Precios referenciales"
COLUMNA_PRODUCTO = "producto original"
COLUMNA_PRECIO = "precio referencial clp"


def normalizar_texto(valor):
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(
        caracter for caracter in texto
        if not unicodedata.combining(caracter)
    ).casefold()
    return re.sub(r"[^a-z0-9]+", " ", texto).strip()


def convertir_precio(valor):
    if valor in (None, ""):
        return None
    if isinstance(valor, bool):
        raise ValueError("El precio no puede ser verdadero/falso.")

    texto = str(valor).strip().replace("$", "").replace(" ", "")
    if not isinstance(valor, (int, float, Decimal)):
        if "." in texto and "," not in texto:
            texto = texto.replace(".", "")
        elif "," in texto and "." not in texto:
            texto = texto.replace(",", ".")

    try:
        precio = Decimal(texto)
    except (InvalidOperation, ValueError) as error:
        raise ValueError(f"Precio inválido: {valor}") from error

    if not precio.is_finite() or precio <= 0:
        raise ValueError("El precio referencial debe ser mayor a cero.")
    return int(precio) if precio == precio.to_integral_value() else float(precio)


def leer_precios_referenciales(ruta):
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        if HOJA_PRECIOS not in libro.sheetnames:
            raise ValueError(
                f"El Excel no contiene la hoja '{HOJA_PRECIOS}'."
            )
        hoja = libro[HOJA_PRECIOS]
        filas = hoja.iter_rows(values_only=True)
        encabezados_originales = next(filas, None)
        if not encabezados_originales:
            raise ValueError("La hoja de precios está vacía.")

        encabezados = {
            normalizar_texto(valor): indice
            for indice, valor in enumerate(encabezados_originales)
            if valor not in (None, "")
        }
        faltantes = [
            columna for columna in (COLUMNA_PRODUCTO, COLUMNA_PRECIO)
            if columna not in encabezados
        ]
        if faltantes:
            raise ValueError(
                "Faltan las columnas requeridas: " + ", ".join(faltantes)
            )

        indice_producto = encabezados[COLUMNA_PRODUCTO]
        indice_precio = encabezados[COLUMNA_PRECIO]
        resultados = []
        sin_precio = 0
        filas_revisadas = 0
        errores = []

        for numero_fila, valores in enumerate(filas, start=2):
            nombre = valores[indice_producto] if len(valores) > indice_producto else None
            if nombre in (None, ""):
                continue
            filas_revisadas += 1
            valor_precio = (
                valores[indice_precio]
                if len(valores) > indice_precio else None
            )
            if valor_precio in (None, ""):
                sin_precio += 1
                continue
            try:
                precio = convertir_precio(valor_precio)
            except ValueError as error:
                errores.append(
                    f"Fila {numero_fila} ({nombre}): {error}"
                )
                continue
            resultados.append({
                "fila": numero_fila,
                "nombre": str(nombre).strip(),
                "precio_referencial": precio,
            })

        nombres = [normalizar_texto(fila["nombre"]) for fila in resultados]
        duplicados = sorted({n for n in nombres if nombres.count(n) > 1})
        if duplicados:
            raise ValueError(
                "El Excel contiene productos repetidos en la carga de precios."
            )
        if errores:
            muestra = "\n".join(errores[:5])
            raise ValueError(
                f"Hay {len(errores)} precios inválidos.\n{muestra}"
            )

        return {
            "filas": resultados,
            "filas_revisadas": filas_revisadas,
            "con_precio": len(resultados),
            "sin_precio": sin_precio,
        }
    finally:
        libro.close()


def preparar_actualizaciones_precios(datos_excel, productos):
    indice = {}
    for producto in productos or []:
        indice.setdefault(
            normalizar_texto(producto.get("nombre")), []
        ).append(producto)

    actualizaciones = []
    sin_coincidencia = []
    ambiguos = []
    ya_valorizados = []
    inactivos = []

    for fila in datos_excel.get("filas", []):
        coincidencias = indice.get(normalizar_texto(fila["nombre"]), [])
        if not coincidencias:
            sin_coincidencia.append(fila)
            continue
        if len(coincidencias) != 1:
            ambiguos.append(fila)
            continue

        producto = coincidencias[0]
        if str(producto.get("estado") or "Activo").casefold() != "activo":
            inactivos.append(fila)
        precio_actual = Decimal(str(producto.get("precio_unitario") or 0))
        if precio_actual != 0:
            ya_valorizados.append({"excel": fila, "producto": producto})
            continue

        actualizaciones.append({
            "producto_id": producto.get("id"),
            "producto": producto.get("nombre"),
            "precio": fila["precio_referencial"],
            "fila_excel": fila["fila"],
        })

    return {
        **datos_excel,
        "actualizaciones": actualizaciones,
        "sin_coincidencia": sin_coincidencia,
        "ambiguos": ambiguos,
        "ya_valorizados": ya_valorizados,
        "inactivos": inactivos,
    }
